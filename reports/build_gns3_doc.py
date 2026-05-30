"""
Converts GNS3_MIGRATION_DECISIONS.md to a formatted Word document.
Same rendering engine as build_docx.py.
"""

import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

MD_PATH  = Path(__file__).parent / "GNS3_MIGRATION_DECISIONS.md"
OUT_PATH = Path(__file__).parent / "GNS3_MIGRATION_DECISIONS.docx"

C_BLACK   = RGBColor(0x1A, 0x1A, 0x1A)
C_WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
C_BLUE    = RGBColor(0x1F, 0x35, 0x64)
C_ACCENT  = RGBColor(0xC0, 0x00, 0x00)
C_CODE_BG = RGBColor(0xF2, 0xF2, 0xF2)
C_TH_BG   = RGBColor(0x1F, 0x35, 0x64)
C_ALT_BG  = RGBColor(0xE9, 0xEF, 0xF8)

HEX_TH_BG  = "1F3564"
HEX_ALT_BG = "E9EFF8"


def set_cell_bg(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_horizontal_rule(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pb = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F3564")
    pb.append(bottom)
    pPr.append(pb)
    p.paragraph_format.space_after = Pt(4)


def apply_inline_styles(run_text: str, para):
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    parts = pattern.split(run_text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            r = para.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = para.add_run(part[1:-1])
            r.font.name = "Courier New"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        elif part:
            para.add_run(part)


def add_code_block(doc, lines: list):
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F2F2F2")
        pPr.append(shd)
        r = p.add_run(line)
        r.font.name = "Courier New"
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(4)


def parse_table(doc, table_lines: list):
    rows = []
    for line in table_lines:
        if re.match(r"^\|[-| :]+\|$", line.strip()):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        rows.append(cells)

    if not rows:
        return

    ncols = len(rows[0])
    tbl = doc.add_table(rows=len(rows), cols=ncols)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl.style = "Table Grid"

    for r_idx, row_data in enumerate(rows):
        row = tbl.rows[r_idx]
        for c_idx, cell_text in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            apply_inline_styles(cell_text, p)
            for run in p.runs:
                run.font.size = Pt(9.5)
            if r_idx == 0:
                set_cell_bg(cell, HEX_TH_BG)
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = C_WHITE
            elif r_idx % 2 == 0:
                set_cell_bg(cell, HEX_ALT_BG)

    doc.add_paragraph()


def build_document(md_text: str) -> Document:
    doc = Document()

    for section in doc.sections:
        section.top_margin    = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = C_BLACK

    lines = md_text.splitlines()
    i = 0

    while i < len(lines):
        raw = lines[i]
        stripped = raw.strip()

        if stripped in ("---", "***", "___"):
            add_horizontal_rule(doc)
            i += 1
            continue

        if stripped.startswith("```"):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i].rstrip())
                i += 1
            add_code_block(doc, code_lines)
            i += 1
            continue

        if stripped.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            parse_table(doc, table_lines)
            continue

        m = re.match(r"^(#{1,4})\s+(.*)", stripped)
        if m:
            level = len(m.group(1))
            text  = m.group(2).strip()
            text  = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)

            if level == 1:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after  = Pt(6)
                r = p.add_run(text)
                r.font.name  = "Calibri"
                r.font.size  = Pt(22)
                r.font.bold  = True
                r.font.color.rgb = C_ACCENT
                p.alignment  = WD_ALIGN_PARAGRAPH.CENTER
            elif level == 2:
                p = doc.add_heading("", level=1)
                p.clear()
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after  = Pt(4)
                r = p.add_run(text)
                r.font.name  = "Calibri"
                r.font.size  = Pt(14)
                r.font.bold  = True
                r.font.color.rgb = C_BLUE
                add_horizontal_rule(doc)
            elif level == 3:
                p = doc.add_heading("", level=2)
                p.clear()
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after  = Pt(2)
                r = p.add_run(text)
                r.font.name  = "Calibri"
                r.font.size  = Pt(12)
                r.font.bold  = True
                r.font.color.rgb = C_BLUE
            else:
                p = doc.add_paragraph()
                r = p.add_run(text)
                r.font.size  = Pt(11)
                r.font.bold  = True
                r.font.color.rgb = C_BLACK
            i += 1
            continue

        m = re.match(r"^(\s*)[-*]\s+(.*)", raw)
        if m:
            indent_level = len(m.group(1)) // 2
            content = m.group(2)
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Cm(0.5 + indent_level * 0.5)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after  = Pt(1)
            p.clear()
            apply_inline_styles(content, p)
            for run in p.runs:
                run.font.size = Pt(10.5)
            i += 1
            continue

        m = re.match(r"^\s*\d+\.\s+(.*)", stripped)
        if m:
            content = m.group(1)
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after  = Pt(1)
            p.clear()
            apply_inline_styles(content, p)
            for run in p.runs:
                run.font.size = Pt(10.5)
            i += 1
            continue

        m = re.match(r"^>\s+(.*)", stripped)
        if m:
            content = m.group(1)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent  = Cm(1.0)
            p.paragraph_format.right_indent = Cm(1.0)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after  = Pt(2)
            pPr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), "FFF3CD")
            pPr.append(shd)
            r = p.add_run("⚠ " + content)
            r.font.size = Pt(10)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x85, 0x66, 0x04)
            i += 1
            continue

        if stripped == "":
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after  = Pt(4)
        apply_inline_styles(stripped, p)
        for run in p.runs:
            run.font.size = Pt(10.5)
        i += 1

    return doc


if __name__ == "__main__":
    md_text = MD_PATH.read_text(encoding="utf-8")
    doc = build_document(md_text)
    doc.save(OUT_PATH)
    print(f"[OK] Saved: {OUT_PATH}")
